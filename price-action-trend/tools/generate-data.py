from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
import argparse
import json

import openpyxl


PRODUCTS = [
    {
        "id": "storage",
        "label": "Storage",
        "color": "#ff8a3d",
        "overall": "Storage Overall",
        "models": ["FS5600", "FS7600", "FS9600"],
    },
    {
        "id": "power",
        "label": "Power Server",
        "color": "#4f8cff",
        "overall": "Power Server Overall",
        "models": ["S1122", "S1124", "E1150", "E1180"],
    },
    {
        "id": "tape",
        "label": "TAPE Library",
        "color": "#b36bff",
        "overall": "TAPE Library Overall",
        "models": [],
    },
]


def date_key(value) -> str:
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    text = str(value).strip().replace("/", "-")
    parts = text.split("-")
    if len(parts) != 3:
        raise ValueError(f"無法辨識日期：{value!r}")
    return f"{int(parts[0]):04d}-{int(parts[1]):02d}-{int(parts[2]):02d}"


def clean(value):
    if isinstance(value, float):
        return round(value, 8)
    return value


def first(record: dict, *names):
    for name in names:
        if name in record:
            return record.get(name)
    return None


def require_sheets(workbook, names):
    missing = [name for name in names if name not in workbook.sheetnames]
    if missing:
        raise ValueError("Excel 缺少工作表：" + "、".join(missing))


def build_payload(workbook_path: Path) -> dict:
    wb = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
    require_sheets(wb, ["圖表資料", "指數總覽", "調價明細", "Overall漲價明細"])

    chart_ws = wb["圖表資料"]
    chart_rows = list(chart_ws.iter_rows(values_only=True))
    if len(chart_rows) < 2:
        raise ValueError("圖表資料沒有可發布的資料列。")
    headers = [str(value).strip() if value is not None else "" for value in chart_rows[0]]
    populated_rows = [row for row in chart_rows[1:] if row[0] is not None]
    dates = [date_key(row[0]) for row in populated_rows]
    if dates != sorted(dates) or len(dates) != len(set(dates)):
        raise ValueError("圖表資料的日期必須依時間遞增且不可重複。")

    series = {}
    for col in range(1, len(headers)):
        name = headers[col]
        if not name:
            continue
        values = [clean(row[col]) if col < len(row) else None for row in populated_rows]
        if any(value is None for value in values):
            raise ValueError(f"圖表資料「{name}」含空白指數。請先讓 Excel 完成重算並儲存。")
        series[name] = values

    expected_series = [item for product in PRODUCTS for item in [product["overall"], *product["models"]]]
    missing_series = [name for name in expected_series if name not in series]
    if missing_series:
        raise ValueError("圖表資料缺少系列：" + "、".join(missing_series))

    summary_ws = wb["指數總覽"]
    event_content = {}
    for col in range(2, summary_ws.max_column + 1):
        raw_date = summary_ws.cell(1, col).value
        if raw_date is not None:
            event_content[date_key(raw_date)] = summary_ws.cell(2, col).value or ""
    if set(event_content) != set(dates):
        raise ValueError("指數總覽與圖表資料的時間節點不一致。")

    overall_ws = wb["Overall漲價明細"]
    overall_headers = [cell.value for cell in next(overall_ws.iter_rows())]
    event_sources = {
        dates[0]: {
            "announcementDate": "",
            "excel": workbook_path.name,
            "pdf": "",
            "note": "基準指數",
        }
    }
    for row in overall_ws.iter_rows(min_row=2, values_only=True):
        record = dict(zip(overall_headers, row))
        effective = record.get("生效日期")
        if not effective:
            continue
        event_sources[date_key(effective)] = {
            "announcementDate": date_key(record["公告日期"]) if record.get("公告日期") else "",
            "excel": record.get("Excel 來源") or "",
            "pdf": record.get("PDF 來源") or "",
            "note": record.get("來源範圍／備註") or "",
        }

    detail_ws = wb["調價明細"]
    detail_headers = [cell.value for cell in next(detail_ws.iter_rows())]
    details = defaultdict(dict)
    for row in detail_ws.iter_rows(min_row=2, values_only=True):
        record = dict(zip(detail_headers, row))
        model = record.get("型號")
        effective = record.get("生效日期")
        if not model or not effective:
            continue
        adopted = clean(first(record, "採用漲幅（公告下限）", "採用平均漲幅", "採用漲幅"))
        details[str(model)][date_key(effective)] = {
            "mtm": record.get("MTM") or "",
            "overall": clean(first(record, "公告參考漲幅", "公告 Overall")),
            "fx": clean(record.get("TWD 匯率")),
            "memory": clean(record.get("記憶體")),
            "processor": clean(record.get("處理器")),
            "ssd": clean(record.get("SSD／Flash")),
            "system": clean(record.get("系統／機體")),
            "cache": clean(record.get("Cache")),
            "tape": clean(record.get("TAPE Library")),
            "other": clean(record.get("其他")),
            "average": adopted,
            "adopted": adopted,
            "combined": clean(record.get("當期綜合漲幅")),
            "before": clean(record.get("調整前指數")),
            "after": clean(record.get("調整後指數")),
            "summary": record.get("漲價內容／差異") or "",
            "status": record.get("資料狀態") or "",
            "excelSource": record.get("Excel 來源") or "",
            "pdfSource": record.get("PDF 來源") or "",
        }

    for series_name in expected_series:
        missing_dates = [key for key in dates if key not in details.get(series_name, {})]
        if missing_dates:
            raise ValueError(f"調價明細「{series_name}」缺少節點：{'、'.join(missing_dates)}")

    modified_at = datetime.fromtimestamp(workbook_path.stat().st_mtime).astimezone()
    return {
        "meta": {
            "title": "IBM Power、Storage、TAPE Library 調價趨勢",
            "baseline": dates[0],
            "baselineIndex": 100,
            "latestDate": dates[-1],
            "source": workbook_path.name,
            "sourceModifiedAt": modified_at.isoformat(timespec="seconds"),
            "generatedAt": datetime.now().astimezone().isoformat(timespec="seconds"),
            "methodVersion": "v2-conservative-floor",
            "method": "同類別多筆先平均；有公告 Overall 時，採公告下限與適用零件類別平均兩者較高值；TWD 匯率獨立累乘；各節點按生效日累乘。",
        },
        "dates": dates,
        "eventContent": event_content,
        "eventSources": event_sources,
        "products": PRODUCTS,
        "series": series,
        "details": details,
    }


def main():
    parser = argparse.ArgumentParser(description="由主 Excel 產生網站資料。")
    parser.add_argument("--workbook", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--json-output", type=Path)
    args = parser.parse_args()

    workbook_path = args.workbook.expanduser().resolve()
    if not workbook_path.exists():
        raise FileNotFoundError(workbook_path)
    payload = build_payload(workbook_path)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    json_text = json.dumps(payload, ensure_ascii=False, indent=2)
    args.output.write_text("window.PRICE_TREND_DATA = " + json_text + ";\n", encoding="utf-8")
    if args.json_output:
        args.json_output.parent.mkdir(parents=True, exist_ok=True)
        args.json_output.write_text(json_text + "\n", encoding="utf-8")
    print(f"資料節點：{len(payload['dates'])}；系列：{len(payload['series'])}")
    print(f"最新節點：{payload['meta']['latestDate']}")
    print(f"已產生：{args.output}")


if __name__ == "__main__":
    main()
